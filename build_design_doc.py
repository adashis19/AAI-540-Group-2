"""
Fill the empty AAI-540 ML Design Document template sections and save a completed
.docx. The original template has section headers whose prompts were left blank
(Data Engineering, Training Data, Feature Engineering, Model Training &
Evaluation, Deployment, Monitoring, CI/CD, Security). This script inserts the
project-specific answers immediately after each such header, preserving the
template structure and the already-written earlier sections.

Run:  python build_design_doc.py
"""
from pathlib import Path
from docx import Document

SRC = Path("/home/ekattim/.kiro/USD_Project/"
           "AAI-540 ML Design Document Template _Group 2.docx")
OUT = Path(__file__).resolve().parent / "reports" / \
    "AAI-540_ML_Design_Document_Group2_COMPLETED.docx"

# Answer blocks keyed by the header text they should follow.
ANSWERS = {
    "Data Engineering:": [
        "Storage: Raw BRFSS SAS-XPT files are kept in a raw store (data/raw/, an "
        "S3 raw bucket in the cloud). Cleaned tabular data is written as columnar "
        "parquet to data/processed/ (an S3 processed/curated bucket, optionally "
        "registered in a feature store). Parquet is used for compression and fast "
        "columnar reads.",
        "Pre-processing before the ML system: strip trailing spaces in XPT column "
        "names; select the harmonised variable set common to 2021/2022/2023; "
        "derive the 0/1 target from _MICHD and drop rows with unknown labels; "
        "convert 'don't know / refused / not asked' codes (e.g. 7/9, 77/99) to "
        "missing values; map 'none' day-codes (88) to 0; rescale BMI (_BMI5/100) "
        "and clip implausible values; tag each row with survey_year.",
    ],
    "Training Data:": [
        "Split: temporal by survey year — train = 2021 + 2022, test = 2023. This "
        "simulates deployment on a future year and is stricter than a random "
        "split. A validation subset can be carved from the training years for "
        "calibration and threshold tuning.",
        "Data labeling: none required. Labels come directly from the "
        "CDC-calculated _MICHD variable, which is authoritative and consistent "
        "across years, so no manual or weak-supervision labeling is used.",
    ],
    "Feature Engineering:": [
        "Fields used (15): bmi, mental_health_days, phys_health_days "
        "(continuous); age_group, sex, gen_health, smoked_100, smoker_status, "
        "diabetes, diff_walk, phys_activity, kidney_disease, income_group, "
        "education_group, has_personal_doctor (categorical/ordinal survey codes).",
        "Fields excluded: _MICHD, CVDINFR4, CVDCRHD4 (target leakage). "
        "Hypertension and cholesterol are excluded because BRFSS collects them "
        "only in odd years and they are absent from the 2022 core file; to keep "
        "one schema valid across all three years they are dropped from the common "
        "feature set (documented limitation). SLEPTIM1 (sleep) appears only in "
        "2022 and is likewise dropped.",
        "Combining / bucketing: age uses the CDC 14-level _AGEG5YR band; income "
        "and education use the grouped CDC variables (_INCOMG1, _EDUCAG); "
        "day-count items are bounded to 0-30; for subgroup analysis age bands are "
        "grouped into 18-39 / 40-59 / 60+.",
        "Other transformations: median imputation + standard scaling for "
        "continuous features; most-frequent imputation + one-hot encoding "
        "(handle_unknown='ignore') for categoricals, all inside a "
        "ColumnTransformer fitted on the training data only to prevent leakage.",
    ],
    "Model Training & Evaluation:": [
        "How trained: each model is combined with the shared preprocessor in a "
        "scikit-learn Pipeline and fitted on the 2021+2022 rows. Class imbalance "
        "is handled with class_weight='balanced' (LogReg, Decision Tree), "
        "balanced_subsample (Random Forest) and scale_pos_weight = neg/pos "
        "(XGBoost).",
        "Algorithms: Logistic Regression (baseline), Decision Tree, Random "
        "Forest, Gradient Boosting, XGBoost; plus a majority-class dummy for "
        "reference.",
        "Key parameters: LogReg C=1.0, max_iter=1000; Decision Tree max_depth=6, "
        "min_samples_leaf=200; Random Forest n_estimators=300, max_depth=12, "
        "min_samples_leaf=50; Gradient Boosting n_estimators=200, max_depth=3, "
        "lr=0.1; XGBoost n_estimators=400, max_depth=5, lr=0.05, subsample=0.9, "
        "colsample_bytree=0.9, tree_method='hist'; global random seed = 42.",
        "Evaluation on the 2023 temporal test set (measured): "
        "MajorityBaseline ROC-AUC 0.500 / PR-AUC 0.085 / recall 0.000; "
        "LogisticRegression 0.834 / 0.315 / 0.802; DecisionTree 0.784 / 0.264 / "
        "0.689; RandomForest 0.828 / 0.311 / 0.790; GradientBoosting 0.835 / "
        "0.319 / 0.079; XGBoost 0.836 / 0.321 / 0.819. XGBoost is selected by "
        "PR-AUC (best for the rare positive class); all learners beat the "
        "majority baseline and class-weighted logistic regression is a close, "
        "interpretable runner-up.",
    ],
    "Model Deployment:": [
        "Instance size: the fitted pipeline is small; a modest CPU instance "
        "(e.g. ml.m5.large, 2 vCPU / 8 GB) scores hundreds of thousands of rows "
        "quickly. No GPU is required.",
        "Batch or real time: batch. Population screening is periodic (over a data "
        "extract), so scheduled batch scoring (e.g. SageMaker Batch Transform or "
        "a containerised job) fits better than a low-latency real-time endpoint "
        "and is cheaper and simpler to operate.",
    ],
    "Model Monitoring:": [
        "Model: track recall, PR-AUC and calibration on newly labelled data and "
        "alert if recall falls below an agreed floor.",
        "Infrastructure: monitor batch-job success/failure, runtime, resource "
        "usage and error rates via CloudWatch-style metrics and alarms.",
        "Data: monitor feature distributions (age, BMI, general health) and "
        "missingness against the training profile (e.g. population-stability "
        "index), and re-compute subgroup recall by sex and age each cycle to "
        "catch fairness drift.",
    ],
    "Model CI/CD:": [
        "Checkpoints: data-schema validation, ingestion/harmonisation, feature "
        "build, training, an evaluation gate, model registration, and "
        "batch-deployment.",
        "Tests: unit tests for ingestion, harmonisation and leakage exclusion; a "
        "data-contract test that fails if a required variable is missing for a "
        "year; an evaluation-gate test that blocks promotion unless PR-AUC and "
        "recall on a fixed hold-out meet or exceed the current production model; "
        "and a smoke test that loads the serialised pipeline and scores a sample "
        "batch.",
    ],
    "Security Checklist, Privacy and Other Risks:": [
        "Store/process PHI? No — BRFSS public files are de-identified survey "
        "responses, not protected health information.",
        "Store/process PII? No direct identifiers (no names, addresses or phone "
        "numbers); only de-identified survey attributes are used.",
        "User behaviour tracked/stored? No. Store/process credit-card "
        "information? No.",
        "Justification: the dataset is a public, de-identified CDC release; the "
        "project neither collects nor links identifiable records.",
        "S3 buckets: a raw bucket (read) for the XPT files and a "
        "processed/curated bucket (read/write) for parquet, model artifacts and "
        "reports, with encryption at rest and least-privilege IAM policies.",
        "Data bias: self-report and recall bias; telephone-survey selection bias; "
        "class imbalance; under-representation of undiagnosed cases in the "
        "negative class.",
        "Bias along sensitive features? Yes — measured subgroup recall: male "
        "0.870 vs female 0.747; age 18-39 0.126, age 40-59 0.589, age 60+ 0.889. "
        "Recall is higher for males and older adults and very low for younger "
        "adults (~1% prevalence); the tool must not be used to rule out risk in "
        "younger adults and these gaps must be monitored.",
        "Ethical concerns: present the model as a screening aid, not a diagnosis; "
        "risk scores could cause anxiety or be misused for discriminatory "
        "decisions (e.g. insurance). Use is restricted to academic prioritisation "
        "with transparency about limitations and subgroup performance.",
    ],
    "Future Enhancements:": [
        "1. Calibration & threshold policy: add isotonic/Platt calibration on a "
        "validation split and select the operating threshold from an explicit "
        "screening budget rather than a fixed 0.5 cut-off.",
        "2. Odd-year enriched model: train a 2021/2023 model that includes blood "
        "pressure and cholesterol and compare it against the harmonised cross-year "
        "model to quantify the value of those predictors.",
        "3. Survey-weighted analysis: incorporate BRFSS complex-survey weights "
        "and report weighted metrics to better reflect the US adult population.",
        "4. Explainability & automated fairness gates: add SHAP-based explanations "
        "and wire subgroup-recall fairness checks into the CI/CD evaluation gate "
        "so regressions block promotion automatically.",
    ],
}


def main():
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    def insert_after(paragraph, text):
        new_p_el = paragraph._p.makeelement(qn("w:p"), {})
        paragraph._p.addnext(new_p_el)
        np = Paragraph(new_p_el, paragraph._parent)
        np.add_run(text)
        return np

    doc = Document(str(SRC))

    # Update authors / instructor to the latest values across all runs.
    REPLACEMENTS = {
        "Authors: Ashis Das, Ram, Hanmant":
            "Authors: Ashis Das, Hanmant Mhalappa Kattimani, Ram Prasanth G",
        "Publication Date: TBD": "Publication Date: October 18, 2026",
    }
    for p in doc.paragraphs:
        for old, new in REPLACEMENTS.items():
            if old in p.text:
                new_text = p.text.replace(old, new)
                # Collapse into the first run, clear the rest.
                if p.runs:
                    p.runs[0].text = new_text
                    for r in p.runs[1:]:
                        r.text = ""
                else:
                    p.add_run(new_text)
    # Add an Instructor line if not present (template had none).
    has_instructor = any("Instructor:" in p.text for p in doc.paragraphs)
    if not has_instructor:
        for p in doc.paragraphs:
            if p.text.strip().startswith("Authors:"):
                from docx.oxml.ns import qn as _qn
                from docx.text.paragraph import Paragraph as _P
                el = p._p.makeelement(_qn("w:p"), {})
                p._p.addnext(el)
                _P(el, p._parent).add_run("Instructor: Dr. Raj Garg")
                break

    paras = {p.text.strip(): p for p in doc.paragraphs
             if p.text.strip() in ANSWERS}
    missing = set(ANSWERS) - set(paras)
    if missing:
        print("WARNING: headers not found:", missing)
    for header, para in paras.items():
        anchor = para
        for block in ANSWERS[header]:
            anchor = insert_after(anchor, block)

    # ----------------------------------------------------------------- #
    # Insert "How We Measure Impact" + "AWS SageMaker Implementation
    # Mapping" immediately BEFORE the "Solution Overview" heading so the
    # docx matches ML_Design_Document_Completed.md.
    # ----------------------------------------------------------------- #
    sol = next((p for p in doc.paragraphs
                if p.text.strip() == "Solution Overview"), None)
    if sol is not None:
        def insert_before(paragraph, text, bold=False, heading=False):
            el = paragraph._p.makeelement(qn("w:p"), {})
            paragraph._p.addprevious(el)
            np = Paragraph(el, paragraph._parent)
            if heading:
                try:
                    np.style = doc.styles["Heading 3"]
                except KeyError:
                    pass
            run = np.add_run(text)
            run.bold = bold or heading
            return np

        impact_blocks = [
            ("How We Measure Impact", True, True),
            ("Impact is measured against the project goals using metrics chosen "
             "for a rare positive class (~8.5% prevalence), so success means "
             "catching true heart-disease cases rather than maximising raw "
             "accuracy. Each goal maps to a concrete, measurable indicator.",
             False, False),
            ("Identify adults likely to report CHD/MI -> minority-class "
             "detection: primary metric recall (0.819), with PR-AUC (0.321) and "
             "ROC-AUC (0.836) on the held-out 2023 year.", False, False),
            ("Beat a naive baseline -> lift: PR-AUC 0.321 vs majority-dummy 0.085 "
             "(~3.8x) and vs logistic-regression baseline 0.315.", False, False),
            ("Produce usable risk scores -> calibration: Brier score and "
             "calibration curve so scores rank and prioritise individuals.",
             False, False),
            ("Evaluate across groups and years -> fairness and temporal "
             "stability: subgroup recall/PR-AUC by sex and age, evaluated on a "
             "future year (2023). Gaps are reported (recall 0.870 male vs 0.747 "
             "female; 0.126 for ages 18-39).", False, False),
            ("Reproducible pipeline -> reliability: passing CI tests "
             "(data-contract, leakage, temporal-split, preprocessor) and an "
             "evaluation gate blocking any model below the incumbent's PR-AUC "
             "and recall.", False, False),
            ("Deployed business-impact proxy: number of true high-risk adults "
             "surfaced per screening cycle at an agreed precision/recall "
             "operating point, tracked alongside subgroup recall to keep the "
             "benefit fair.", False, False),
            ("AWS SageMaker Implementation Mapping", True, True),
            ("The system is designed to run on AWS SageMaker; each pipeline stage "
             "maps to a managed service, and this mapping is what the video "
             "demonstration screencasts. The local scikit-learn implementation "
             "in src/ is the reference logic the SageMaker components execute.",
             False, False),
            ("S3 raw and processed buckets -> versioned XPT and parquet objects. "
             "SageMaker Feature Store -> the 15-feature group schema and record "
             "ingestion. SageMaker Training Jobs -> training logs and metrics. "
             "SageMaker Model Registry -> registered model package group and the "
             "approved version. SageMaker Pipelines (DAG) -> a successful run and "
             "an intentionally failed run at the evaluation gate. SageMaker Batch "
             "Transform -> the scored-records output. SageMaker Model Monitor and "
             "CloudWatch -> data-quality/model-quality reports and infrastructure "
             "dashboards.", False, False),
        ]
        for text, bold, heading in impact_blocks:
            insert_before(sol, text, bold=bold, heading=heading)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print("Wrote", OUT, "| filled headers:", list(paras))


if __name__ == "__main__":
    main()
